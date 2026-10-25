#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
祐興工作台 API 後端
提供數據讀取和狀態檢查功能
"""
import os
import json
import re
from datetime import datetime
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs
import openpyxl

# 文件路徑
# 2026-10-23 改：CDM live 真身已搬去 C:\CDM（脫離 OneDrive）。
#   INVENTORY_FILE 保持 Claw 副本（倉存 data 只喺嗰度，含 9 萬補貨）
#   CDM_DIR 改去 live C:\CDM，令 live API（/api/cdm-status）同 status.json 快照一致
INVENTORY_FILE = r'C:\Users\benso\WorkBuddy\Claw\cdm\膠盒倉存(WB).xlsx'
CDM_DIR = r'C:\CDM'
# 發票總額來源：CDM live 真身嘅 Flow1 單日發票（唯一真相源）
INVOICE_BASE = r'C:\CDM\CDM 發票對數\CDM 發票'
INV_RE = re.compile(r'^inv_(\d{2})-(\d{2})-(\d{4})\.xlsx$')

# 🚨 2026-10-03 修：倉存頁結餘一定要用「當月 sheet + 最新有日期嗰行 + col6/col7」
#    - 舊 code 用 wb.active（openpyxl 開出嚟第一頁 = 「4月」）→ 回 4 月嘅數字
#    - 舊 code 又用 col4/col5（標準頁 D/E 係「補貨」欄，唔係結餘）→ 出 '-'
# 標準頁欄位：A日期 B/C用量 D/E補貨 F/G結餘 H總用量
DATE_RE = re.compile(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}')


def read_box_latest(path=INVENTORY_FILE):
    """回 {'small':int,'big':int,'date':str} —— 當月 sheet 最新有日期嗰行嘅結餘"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        today = datetime.now()
        sheet = None
        for name in wb.sheetnames:                      # 1) 揀當月 sheet
            if str(name) == '%d月' % today.month:
                sheet = wb[name]
                break
        if sheet is None:                                # 2) 冇就 fallback 翻最新嘅月份頁
            sheet = wb[wb.sheetnames[-1]]

        last = None                                      # 3) 搵最後一個「有日期」嘅行
        for r in range(sheet.max_row, 0, -1):
            v = sheet.cell(r, 1).value
            if v is None:
                continue
            if DATE_RE.match(str(v).strip()):
                last = r
                break
        if last is None:
            return {'small': 0, 'big': 0, 'date': '未知'}

        small = sheet.cell(last, 6).value or 0           # F = 小膠盒結餘
        big = sheet.cell(last, 7).value or 0             # G = 大膠盒結餘
        date = sheet.cell(last, 1).value
        return {'small': small, 'big': big,
                'date': str(date)[:10] if date else '未知'}
    finally:
        wb.close()


def read_box_history(month=None):
    """回 {success, month, months, rows:[{date,small_bal,big_bal,small_used,big_used}], updated_at}
       讀指定月份 sheet 嘅每日結餘/用量。欄位用 header(row2) 配對，兼容 8 欄新格式同 5 欄舊格式。"""
    wb = openpyxl.load_workbook(INVENTORY_FILE, data_only=True)
    try:
        months = list(wb.sheetnames)
        sheet = None
        if month and month in months:
            sheet = wb[month]
        else:
            today = datetime.now()
            name = '%d月' % today.month
            sheet = wb[name] if name in months else wb[months[-1]]
            month = sheet.title

        # 用 header(row2) 配對欄位（兼容 5 欄舊 / 8 欄新）
        hdr = {}
        for c in range(1, sheet.max_column + 1):
            v = sheet.cell(2, c).value
            if v is None:
                continue
            s = str(v).strip()
            if '日期' in s:
                hdr.setdefault('date', c)
            elif '1-10' in s and '結餘' in s:
                hdr.setdefault('small_bal', c)
            elif '11-12' in s and '結餘' in s:
                hdr.setdefault('big_bal', c)
            elif '1-10' in s and '用量' in s:
                hdr.setdefault('small_used', c)
            elif '11-12' in s and '用量' in s:
                hdr.setdefault('big_used', c)
        date_c = hdr.get('date', 1)
        small_bal_c = hdr.get('small_bal', 4 if sheet.max_column <= 5 else 6)
        big_bal_c = hdr.get('big_bal', 5 if sheet.max_column <= 5 else 7)
        small_used_c = hdr.get('small_used', 2)
        big_used_c = hdr.get('big_used', 3)

        rows = []
        for r in range(3, sheet.max_row + 1):
            v = sheet.cell(r, date_c).value
            if v is None:
                continue
            if not DATE_RE.match(str(v).strip()):
                continue  # 跳過 總用量 / 期末結餘 等 summary 行
            def _int(x):
                return int(x) if isinstance(x, (int, float)) else 0
            rows.append({
                'date': str(v)[:10],
                'small_bal': _int(sheet.cell(r, small_bal_c).value),
                'big_bal': _int(sheet.cell(r, big_bal_c).value),
                'small_used': _int(sheet.cell(r, small_used_c).value),
                'big_used': _int(sheet.cell(r, big_used_c).value),
            })
        return {'month': month, 'months': months, 'rows': rows,
                'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    finally:
        wb.close()


def _inv_day_total(path):
    """計單張 inv_DD-MM-YYYY.xlsx 嘅發票總額 = Σ(數量×單價)，回 (total, items)。
       欄位：col2 貨號、col13 數量、col17 單價（2026-10-25 實測 inv_26 確認）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        tot = 0.0
        items = 0
        for r in range(1, ws.max_row + 1):
            code = ws.cell(r, 2).value
            qty = ws.cell(r, 13).value
            price = ws.cell(r, 17).value
            if isinstance(code, str) and code.strip() \
                    and isinstance(qty, (int, float)) and isinstance(price, (int, float)):
                tot += qty * price
                items += 1
        return round(tot, 2), items
    finally:
        wb.close()


def read_invoice_history(month=None):
    """回 {success, month, months, rows:[{date,total,items}], updated_at}
       掃 INVOICE_BASE 下 {month} 文件夾內 inv_DD-MM-YYYY.xlsx，計每日發票總額(HK$)。
       month=None 時列晒有 inv 發票嘅月份文件夾，並優先揀當月。"""
    if not os.path.isdir(INVOICE_BASE):
        return {'success': False, 'error': '發票目錄不存在', 'months': [], 'rows': []}
    # 1) 列月份文件夾（含 inv_ 發票者）
    month_dirs = []
    for name in sorted(os.listdir(INVOICE_BASE)):
        d = os.path.join(INVOICE_BASE, name)
        if not os.path.isdir(d):
            continue
        try:
            if any(INV_RE.match(f) for f in os.listdir(d)):
                month_dirs.append(name)
        except OSError:
            continue
    if not month_dirs:
        return {'success': False, 'error': '無發票月份資料', 'months': [], 'rows': []}
    # 2) 選月
    if month and month in month_dirs:
        cur = month
    else:
        today = datetime.now()
        abbr = today.strftime('%b')
        cur = None
        for m in month_dirs:
            if abbr in m and str(today.year) in m:
                cur = m
                break
        if cur is None:
            cur = month_dirs[-1]
    # 3) 讀該月所有 inv 發票
    d = os.path.join(INVOICE_BASE, cur)
    rows = []
    for f in sorted(os.listdir(d)):
        m = INV_RE.match(f)
        if not m or '.bak' in f:
            continue
        dd, mm, yyyy = m.groups()
        date_str = '%s-%s-%s' % (yyyy, mm, dd)
        try:
            total, items = _inv_day_total(os.path.join(d, f))
        except Exception:
            total, items = 0.0, 0
        rows.append({'date': date_str, 'total': total, 'items': items})
    rows.sort(key=lambda x: x['date'])
    return {'success': True, 'month': cur, 'months': month_dirs, 'rows': rows,
            'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')}


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == '/api/inventory':
            self.send_inventory()
        elif self.path.startswith('/api/box_history'):
            self.send_box_history()
        elif self.path.startswith('/api/invoice_history'):
            self.send_invoice_history()
        elif self.path == '/api/cdm-status':
            self.send_cdm_status()
        elif self.path == '/api/check-all':
            self.send_all_check()
        else:
            self.send_error(404, 'Not found')
    
    def send_inventory(self):
        """讀取膠盒倉存數據（當月 sheet 最新日嘅結餘）"""
        try:
            box = read_box_latest()
            data = {
                'success': True,
                'small_box': box['small'],
                'big_box': box['big'],
                'last_update': box['date'],
                'file_path': INVENTORY_FILE
            }
            self.send_json(data)
        except Exception as e:
            self.send_json({'success': False, 'error': str(e)})

    def send_box_history(self):
        """讀取膠盒倉存整月每日資料（?month=10月）"""
        try:
            q = parse_qs(urlparse(self.path).query)
            month = q.get('month', [None])[0]
            data = read_box_history(month)
            data['success'] = True
            self.send_json(data)
        except Exception as e:
            self.send_json({'success': False, 'error': str(e)})

    def send_invoice_history(self):
        """讀取 CDM 每日發票總額（?month=10 Oct 2026）"""
        try:
            q = parse_qs(urlparse(self.path).query)
            month = q.get('month', [None])[0]
            data = read_invoice_history(month)
            self.send_json(data)
        except Exception as e:
            self.send_json({'success': False, 'error': str(e)})
    
    def send_cdm_status(self):
        """檢查 CDM 流程狀態"""
        try:
            # 檢查 CDM 輸出路徑
            output_dir = os.path.join(CDM_DIR, '輸出')
            pdf_files = []
            if os.path.exists(output_dir):
                pdf_files = [f for f in os.listdir(output_dir) if f.endswith('.pdf')]
            
            # 檢查最新的訂單處理記錄
            processed_files = []
            for f in os.listdir(CDM_DIR):
                if f.startswith('processed_') and f.endswith('.json'):
                    processed_files.append(f)
            
            data = {
                'success': True,
                'cdm_running': len(pdf_files) > 0,
                'latest_pdf_count': len(pdf_files),
                'processed_orders': len(processed_files),
                'last_run': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            }
            self.send_json(data)
        except Exception as e:
            self.send_json({'success': False, 'error': str(e)})
    
    def send_all_check(self):
        """執行全盤檢查"""
        result = {
            'success': True,
            'timestamp': datetime.now().isoformat(),
            'inventory': None,
            'cdm': None
        }
        
        # 檢查膠盒倉存
        try:
            box = read_box_latest()
            result['inventory'] = {
                'small_box': box['small'],
                'big_box': box['big'],
                'last_date': box['date']
            }
        except Exception as e:
            result['inventory_error'] = str(e)
        
        # 檢查 CDM
        try:
            output_dir = os.path.join(CDM_DIR, '輸出')
            pdf_files = []
            if os.path.exists(output_dir):
                pdf_files = [f for f in os.listdir(output_dir) if f.endswith('.pdf')]
            result['cdm'] = {
                'running': len(pdf_files) > 0,
                'pdf_count': len(pdf_files)
            }
        except Exception as e:
            result['cdm_error'] = str(e)
        
        self.send_json(result)
    
    def send_json(self, data):
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False).encode('utf-8'))
    
    def send_error(self, code, message):
        self.send_response(code)
        self.send_header('Content-Type', 'application/json')
        self.end_headers()
        self.wfile.write(json.dumps({'success': False, 'error': message}).encode('utf-8'))
    
    def log_message(self, format, *args):
        pass  # 靜默日誌

if __name__ == '__main__':
    port = 8083
    server = HTTPServer(('0.0.0.0', port), Handler)
    print(f'API server running on port {port}')
    server.serve_forever()
